#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
将 PRD.md 转换为 PRD.docx Word 文档
读取 markdown，解析格式（标题/表格/列表/加粗/代码块），应用 Scope 校正，输出 docx
"""

import re
import os
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def apply_scope_corrections(text):
    """
    对 markdown 文本应用 Scope 校正（直接文本替换）
    参数：
        text (str): 原始 markdown 文本
    返回：
        str: 校正后的 markdown 文本
    """
    # === 1. 信誉体系/信誉分/动态衰减 → 信息完整度三档 + 基础实名权限两档 ===
    credibility_replacement = "信息完整度三档🟢🟡🔴 + 基础实名权限两档"
    # 按长度从长到短替换，避免子串提前匹配
    credibility_terms = [
        "信誉评价闭环",
        "评价信誉体系",
        "信誉体系",
        "信誉分",
        "动态衰减",
    ]
    for term in credibility_terms:
        text = text.replace(term, credibility_replacement)

    # === 2. 撮合确认码/撮合反馈/周度自适应调权 → 7天一键刷新 + 信息完整度升级引导 ===
    matching_replacement = "7天一键刷新有效期 + 信息完整度升级引导"
    matching_terms = [
        "撮合确认码",
        "撮合反馈",
        "周度自适应调权",
    ]
    for term in matching_terms:
        text = text.replace(term, matching_replacement)

    # === 3. 站内IM/轻量IM → 联系方式中转页（永久单轨，不接IM SDK） ===
    im_replacement = "联系方式中转页（永久单轨，不接IM SDK）"
    im_terms = [
        "站内 IM",
        "站内IM",
        "轻量 IM",
        "轻量IM",
    ]
    for term in im_terms:
        text = text.replace(term, im_replacement)

    # === 4. QS五档/钻石/金牌/银牌 → 信息完整度三档 ===
    qs_replacement = "信息完整度三档🟢🟡🔴"
    qs_terms = [
        "QS五档",
        "钻石/金牌/银牌",
    ]
    for term in qs_terms:
        text = text.replace(term, qs_replacement)

    # === 5. 北极星指标：近场撮合完成率 → 3步找到率 ===
    # 先替换完整定义行
    text = text.replace(
        "> **近场撮合完成率** = 发布资源/需求后 24h 内获得至少 1 次有效联系的比例",
        "> **3步找到率** = AI匹配召回率 × 图层加载成功率(P95≤300ms) × 🟢完整发布率"
    )
    # 再替换其余出现的"近场撮合完成率"和"撮合完成率"（注意顺序，先长后短）
    text = text.replace("近场撮合完成率", "3步找到率")
    text = text.replace("撮合完成率", "3步找到率")

    # === 6. 11个算法群保留砍除状态，加入 D5.1 AI三色清单 ===
    d5_1_note = "\n- **D5.1 AI三色清单**：绿=LR+GBDT/3B小模型；红=11算法群/70B永不做\n"
    score_block_end = "  （默认权重 w_dist=0.45, w_fresh=0.25, w_match=0.2, w_key=0.05, w_cert=0.05；后台可配）\n  ```\n"
    if score_block_end in text:
        text = text.replace(score_block_end, score_block_end + d5_1_note)

    # === 7. 文档版本、日期、状态 ===
    text = text.replace("| 文档版本 | v2.0 |", "| 文档版本 | v2.1 |")
    text = text.replace("| 文档状态 | 草案（待评审） |", "| 文档状态 | Scope校正版（待评审） |")
    text = text.replace("| 编制日期 | 2026-08-10 |", "| 编制日期 | 2026-08-19 |")

    return text


def add_runs_with_formatting(paragraph, text):
    """
    解析文本中的 **加粗** 和 `代码` 标记，添加对应格式的 run
    参数：
        paragraph: Word 段落对象
        text (str): 含 markdown 格式标记的文本
    """
    # 使用正则分割出 **bold** 和 `code` 部分
    parts = re.split(r'(\*\*.*?\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            # 加粗文本
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            # 等宽代码字体
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def add_heading(doc, text, level):
    """
    添加标题段落（Heading 1-4）
    参数：
        doc: Word 文档对象
        text (str): 标题文本
        level (int): 标题级别（1-4）
    """
    level = max(1, min(level, 4))
    try:
        p = doc.add_heading(level=level)
    except KeyError:
        # 降级：用普通段落 + 粗体大字
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(18 - level * 2)
        return
    add_runs_with_formatting(p, text)


def add_paragraph(doc, text):
    """
    添加普通段落
    参数：
        doc: Word 文档对象
        text (str): 段落文本
    """
    p = doc.add_paragraph()
    add_runs_with_formatting(p, text)


def add_quote(doc, quote_lines):
    """
    添加引用块段落（带左缩进）
    参数：
        doc: Word 文档对象
        quote_lines (list): 引用文本行列表
    """
    for line in quote_lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        add_runs_with_formatting(p, line)


def add_code_block(doc, code_text):
    """
    添加代码块（Consolas 等宽字体，每行一个段落）
    参数：
        doc: Word 文档对象
        code_text (str): 代码块文本
    """
    lines = code_text.split('\n')
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(line if line else ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(9)


def add_list_item(doc, text, bullet=True, indent_level=0):
    """
    添加列表项（使用 Word 内置列表样式）
    参数：
        doc: Word 文档对象
        text (str): 列表项文本
        bullet (bool): True=无序列表，False=有序列表
        indent_level (int): 缩进级别（0=顶级）
    """
    if bullet:
        base_style = 'List Bullet'
    else:
        base_style = 'List Number'

    # 根据缩进级别选择对应样式（List Bullet / List Bullet 2 / List Bullet 3）
    if indent_level == 0:
        style_name = base_style
    else:
        style_name = '%s %d' % (base_style, min(indent_level + 1, 3))

    try:
        p = doc.add_paragraph(style=style_name)
    except KeyError:
        # 样式不存在时降级为普通段落 + 手动缩进 + 前缀符号
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75 + 0.5 * indent_level)
        prefix = '• ' if bullet else '%d. ' % (indent_level + 1)
        run = p.add_run(prefix)
        run.bold = True

    add_runs_with_formatting(p, text)


def add_table(doc, table_lines):
    """
    解析 markdown 表格行并添加 Word 表格（带 Table Grid 边框样式）
    参数：
        doc: Word 文档对象
        table_lines (list): markdown 表格行列表
    """
    # 解析所有行
    rows = []
    for line in table_lines:
        line = line.strip()
        # 去掉首尾的 |
        if line.startswith('|'):
            line = line[1:]
        if line.endswith('|'):
            line = line[:-1]
        cells = [c.strip() for c in line.split('|')]
        rows.append(cells)

    # 检测并跳过分隔行（|---|---| 或 |:---:|）
    body_start = 1
    if len(rows) > 1:
        second_row = rows[1]
        if all(re.match(r'^[-:]+$', c) for c in second_row if c):
            body_start = 2

    header = rows[0] if rows else []
    body = rows[body_start:]

    num_cols = len(header)
    if num_cols == 0:
        return

    # 创建表格并应用 Table Grid 样式（自带边框）
    table = doc.add_table(rows=1 + len(body), cols=num_cols)
    try:
        table.style = 'Table Grid'
    except KeyError:
        pass

    # 填充表头（加粗）
    for j, cell_text in enumerate(header):
        if j < num_cols:
            cell = table.rows[0].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            add_runs_with_formatting(p, cell_text)
            for run in p.runs:
                run.bold = True

    # 填充表体
    for i, row in enumerate(body):
        for j, cell_text in enumerate(row):
            if j < num_cols:
                cell = table.rows[i + 1].cells[j]
                cell.text = ''
                p = cell.paragraphs[0]
                add_runs_with_formatting(p, cell_text)

    # 表格后加空行以分隔后续内容
    doc.add_paragraph()


def setup_document():
    """
    创建并配置 Word 文档（A4 页面、正常页边距、默认字体）
    返回：
        Document: 配置好的 Word 文档对象
    """
    doc = Document()

    # 页面设置：A4 + 正常页边距（2.54cm）
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)

    # 设置默认字体（西文 Calibri，东亚 Microsoft YaHei）
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    return doc


def md_to_docx(md_text, doc):
    """
    将 markdown 文本转换为 Word 文档内容（逐行解析）
    参数：
        md_text (str): markdown 文本
        doc: Word 文档对象
    """
    lines = md_text.split('\n')
    i = 0
    total = len(lines)

    while i < total:
        line = lines[i]

        # --- 代码块 ---
        if line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < total and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, '\n'.join(code_lines))
            i += 1
            continue

        # --- 标题 ---
        heading_match = re.match(r'^(#{1,4})\s+(.*)', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            add_heading(doc, text, level)
            i += 1
            continue

        # --- 表格 ---
        stripped = line.strip()
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < total and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, table_lines)
            continue

        # --- 水平分割线 ---
        if stripped == '---':
            i += 1
            continue

        # --- 引用块 ---
        if stripped.startswith('>'):
            quote_lines = []
            while i < total and lines[i].strip().startswith('>'):
                quote_text = lines[i].strip()[1:].strip()
                quote_lines.append(quote_text)
                i += 1
            add_quote(doc, quote_lines)
            continue

        # --- 无序列表 ---
        bullet_match = re.match(r'^(\s*)[-*]\s+(.*)', line)
        if bullet_match:
            indent_str = bullet_match.group(1)
            text = bullet_match.group(2)
            indent_level = len(indent_str) // 2
            add_list_item(doc, text, bullet=True, indent_level=indent_level)
            i += 1
            continue

        # --- 有序列表 ---
        num_match = re.match(r'^(\s*)(\d+)\.\s+(.*)', line)
        if num_match:
            indent_str = num_match.group(1)
            text = num_match.group(3)
            indent_level = len(indent_str) // 2
            add_list_item(doc, text, bullet=False, indent_level=indent_level)
            i += 1
            continue

        # --- 普通段落 ---
        if stripped:
            add_paragraph(doc, stripped)

        i += 1


def main():
    """
    主函数：读取 PRD.md → 应用 Scope 校正 → 转换为 Word → 保存 PRD.docx → 验证
    """
    base_dir = os.path.dirname(os.path.abspath(__file__))
    prd_md_path = os.path.join(base_dir, 'docs', 'PRD.md')
    prd_docx_path = os.path.join(base_dir, 'docs', 'PRD.docx')

    # 1. 读取 markdown
    with open(prd_md_path, 'r', encoding='utf-8') as f:
        md_content = f.read()
    print('[1/4] 已读取 PRD.md：%d 字符' % len(md_content))

    # 2. 应用 Scope 校正
    md_content = apply_scope_corrections(md_content)
    print('[2/4] 已应用 Scope 校正')

    # 3. 转换为 Word 文档
    doc = setup_document()
    md_to_docx(md_content, doc)
    print('[3/4] 已转换为 Word 格式')

    # 4. 保存并验证
    doc.save(prd_docx_path)
    file_size = os.path.getsize(prd_docx_path)
    print('[4/4] 已保存 PRD.docx：%s' % prd_docx_path)
    print('      文件大小：%d 字节（%.1f KB）' % (file_size, file_size / 1024.0))

    if file_size > 0:
        print('✅ 验证通过：文件生成成功且大小 > 0')
        return 0
    else:
        print('❌ 验证失败：文件大小为 0')
        return 1


if __name__ == '__main__':
    sys.exit(main())
