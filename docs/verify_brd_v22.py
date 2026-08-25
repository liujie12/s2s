"""校验 BRD v2.2 排期基线修订是否已落盘到 BRD.pptx 与 BRD.md。

设计依据：docs/solutions/workflow-issues/multi-format-deliverable-consistency.md
  - 特征句必须用改写后的**原句片段**，不用通用词表（通用词表无法区分「该出现」与「不该出现」）；
  - pptx 提取必须递归展开 Group，否则大量漏检；
  - 同时做反向校验：旧口径（月 1-2 等相对月份）必须已从 pptx 中清零。

用法：python verify_brd_v22.py
返回：无返回值；打印 [OK]/[MISS]/[BAD] 逐条结论与汇总计数，非零缺陷时以退出码 1 结束。
"""

import sys
from pptx import Presentation


PPTX = r"d:\developer\code\aicoding\s2s\docs\找鸭找-商业需求文档-v2.1.pptx"
MD = r"d:\developer\code\aicoding\s2s\docs\BRD.md"

# 本轮在 gen_brd_pptx.py 中改写过的原句片段（slide 20 路线图 + slide 22 资源依赖）
PPTX_MUST = [
    "9/30 交内测包",
    "第一批｜08/24-09/30",
    "五大 P0 闭环",
    "合规窗口｜10/01-11/25",
    "运营治理后台 7 页、公安备案",
    "后续批次｜11/26 起",
    "2027/04 联动",
    "首批为单人加 AI 辅助",
    "软著 60 工作日、ICP 与 APP 备案、企业主体",
]

# 旧相对月份口径必须已从 pptx 清零（v2.2 改为绝对日期锚点）
PPTX_MUST_NOT = ["月 1-2", "月 3-4", "月 5-7", "月 7 三目标", "五人核心团队"]

# BRD.md 侧本轮新增的原句片段
MD_MUST = [
    "v2.2 修订：相对月份 → 绝对日期锚点",
    "2026-08-24 → 09-30",
    "上架合规窗口",
    "B6-v2.2 裁剪说明",
    "二次裁剪触发条件",
    "验收节点归属（v2.2 新增）",
    "v2.2 实际人力说明",
    "上架资质前置项（v2.2 新增",
    "不可为增收而破",
    "9/30 交付形态：内测包而非商店上架",
    "Batch1 范围裁剪：S4 运营治理后台移出",
]


def pptx_text(path):
    """提取 pptx 全文，递归展开 Group 后拼接。

    参数：path —— pptx 文件绝对路径
    返回：str，全部形状文本以换行拼接
    """
    prs = Presentation(path)
    parts = []

    def walk(shapes):
        for shp in shapes:
            if shp.shape_type is not None and str(shp.shape_type).startswith("GROUP"):
                walk(shp.shapes)
                continue
            if shp.has_text_frame:
                parts.append(shp.text_frame.text)

    for slide in prs.slides:
        walk(slide.shapes)
    return "\n".join(parts)


def main():
    """主流程：对 pptx 做正反双向校验、对 md 做正向校验，并汇总缺陷数。"""
    bad = 0
    ptxt = pptx_text(PPTX)
    with open(MD, encoding="utf-8") as f:
        mtxt = f.read()

    print("=== PPTX 正向（必须命中）===")
    for s in PPTX_MUST:
        ok = s in ptxt
        print(f"{'[OK]  ' if ok else '[MISS]'} {s}")
        bad += 0 if ok else 1

    print("=== PPTX 反向（必须清零）===")
    for s in PPTX_MUST_NOT:
        gone = s not in ptxt
        print(f"{'[OK]  ' if gone else '[BAD] '} {s}")
        bad += 0 if gone else 1

    print("=== BRD.md 正向（必须命中）===")
    for s in MD_MUST:
        ok = s in mtxt
        print(f"{'[OK]  ' if ok else '[MISS]'} {s}")
        bad += 0 if ok else 1

    print(f"\ndefects={bad}")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
