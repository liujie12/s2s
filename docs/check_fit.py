"""按形状几何尺寸估算 PPTX 文本是否溢出（文字重叠）风险。

功能：遍历指定 pptx 的全部形状（含嵌套 Group），对含文字的形状按
      「框宽 / 字宽」估算每行可容纳字符数，再按「框高 / 行高」估算可容纳行数，
      与实际文本所需行数比较，超出即报警。
参数：命令行 argv[1] = pptx 路径；argv[2] = 可选，只审包含该关键词的文本
返回：进程退出码 0；结果打印到 stdout
"""
import sys
from pptx import Presentation
from pptx.util import Emu

EMU_PER_PT = 12700


def default_font_size(shape):
    """取形状首个非空 run 的字号（pt），取不到时按 12pt 兜底。

    参数：shape —— python-pptx 形状对象
    返回：float，字号磅值
    """
    for para in shape.text_frame.paragraphs:
        if para.font.size is not None:
            return para.font.size.pt
        for run in para.runs:
            if run.font.size is not None:
                return run.font.size.pt
    return 12.0


def walk(shapes, prefix, out):
    """递归收集形状路径与文本框对象。

    参数：shapes 形状集合；prefix 路径前缀；out 结果列表（原地追加）
    返回：无
    """
    for shp in shapes:
        cur = f"{prefix} > {shp.name}" if prefix else shp.name
        if shp.shape_type is not None and str(shp.shape_type).startswith("GROUP"):
            walk(shp.shapes, cur, out)
            continue
        if shp.has_text_frame and shp.text_frame.text.strip():
            out.append((cur, shp))


def main():
    """主流程：逐页估算并打印溢出风险清单。"""
    path = sys.argv[1]
    keyword = sys.argv[2] if len(sys.argv) > 2 else None
    prs = Presentation(path)
    risky = 0
    for idx, slide in enumerate(prs.slides, 1):
        items = []
        walk(slide.shapes, "", items)
        for cur, shp in items:
            text = shp.text_frame.text.strip()
            if keyword and keyword not in text:
                continue
            if shp.width is None or shp.height is None:
                continue
            size = default_font_size(shp)
            # 中文全角字符宽约等于字号，英文数字约 0.55 倍，按混排折算平均宽
            han = sum(1 for c in text if ord(c) > 0x2E80)
            other = len(text) - han
            avg_w = (han * 1.0 + other * 0.55) / max(len(text), 1)
            char_w_emu = size * avg_w * EMU_PER_PT
            line_h_emu = size * 1.25 * EMU_PER_PT
            # 左右内边距按默认 0.1 英寸各一侧扣除
            usable_w = shp.width - Emu(91440) * 2
            usable_h = shp.height - Emu(45720) * 2
            if usable_w <= 0 or usable_h <= 0:
                continue
            per_line = max(int(usable_w / char_w_emu), 1)
            need_lines = -(-len(text) // per_line)
            cap_lines = max(int(usable_h / line_h_emu), 1)
            if need_lines > cap_lines:
                risky += 1
                print(
                    f"[溢出] S{idx} {cur} size={size} "
                    f"len={len(text)} perline={per_line} "
                    f"need={need_lines} cap={cap_lines} :: {text[:40]}"
                )
    print(f"risky={risky}")


if __name__ == "__main__":
    main()
