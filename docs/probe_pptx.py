"""探测 PPTX 模板结构（含 Group 递归），输出到文件便于分析。

功能：读取指定 pptx，递归展开所有 Group，把每个可写文本的定位路径
      （slide 序号 / 形状名路径 / 段落序号）与当前文字、字号写入文本文件。
参数：
    argv[1] —— 待探测的 pptx 路径
    argv[2] —— 结构报告输出路径（可选，默认 probe_out.txt）
返回：无返回值，结果写入输出文件
"""
import sys
import os

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".pylibs"))

from pptx import Presentation

OUT = []


def walk(shapes, depth=0, path=""):
    """递归遍历形状集合，收集文本信息到全局 OUT 列表。

    参数：
        shapes: 形状集合（slide.shapes 或 group.shapes）
        depth:  递归深度，用于缩进展示层级
        path:   形状名路径，如 "Group 1 > Group 15"
    返回：无
    """
    for shp in shapes:
        pad = "  " * depth
        cur = f"{path} > {shp.name}" if path else shp.name
        is_group = shp.shape_type is not None and str(shp.shape_type).startswith("GROUP")
        if is_group:
            OUT.append(f"{pad}[G] {shp.name}")
            walk(shp.shapes, depth + 1, cur)
            continue
        if shp.has_text_frame:
            texts = []
            for pi, para in enumerate(shp.text_frame.paragraphs):
                t = "".join(r.text for r in para.runs)
                if t.strip():
                    sz = next((r.font.size.pt for r in para.runs if r.font.size), None)
                    texts.append(f"p{pi}({sz}): {t}")
            if texts:
                OUT.append(f"{pad}[T] {cur}")
                for t in texts:
                    OUT.append(f"{pad}     {t}")


def main():
    """主流程：遍历全部幻灯片并写出结构报告。"""
    path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        os.path.dirname(__file__), "probe_out.txt"
    )
    prs = Presentation(path)
    OUT.append(f"slides={len(prs.slides)} layouts={[l.name for l in prs.slide_layouts]}")
    OUT.append(f"size={prs.slide_width.pt}x{prs.slide_height.pt}pt")
    for si, slide in enumerate(prs.slides, 1):
        OUT.append("")
        OUT.append(f"===== Slide {si} | layout={slide.slide_layout.name} =====")
        walk(slide.shapes)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(OUT))
    print(f"written: {out_path} lines={len(OUT)}")


if __name__ == "__main__":
    main()
